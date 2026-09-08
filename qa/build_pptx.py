# -*- coding: utf-8 -*-
"""Build an editable PPTX from the extracted deck op-list.

Geometry is 1:1 with the HTML deck (1920x1080 px -> 13.333x7.5 in, 6350 EMU/px,
1 px = 0.5 pt).  Text stays text, tables stay tables, shapes stay shapes.
"""
import json, os, sys
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from copy import deepcopy

SRC = sys.argv[1]
DST = sys.argv[2]
PX = 6350                      # EMU per deck pixel
FONT = 'Arial'                 # metric-stable Hebrew+Latin font on Windows/Mac
LRE, PDF_ = '‪', '‬'  # LTR embedding for Latin inside RTL text
NOSTYLE = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'  # "No Style, No Grid"

data = json.load(open(SRC + '/ops.json', encoding='utf-8'))
prs = Presentation()
prs.slide_width, prs.slide_height = Emu(1920 * PX), Emu(1080 * PX)
blank = prs.slide_layouts[6]

ALIGN = {'l': PP_ALIGN.LEFT, 'r': PP_ALIGN.RIGHT, 'ctr': PP_ALIGN.CENTER}
ANCH = {'t': MSO_ANCHOR.TOP, 'ctr': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}


def E(v):
    return Emu(int(round(v * PX)))


def style_run(r, spec, rtl, scale=1.0):
    """Apply size/weight/colour + latin & complex-script fonts + bidi isolation."""
    txt = spec['txt']
    if spec.get('ltr') and rtl:
        txt = LRE + txt + PDF_
    r.text = txt
    rPr = r._r.get_or_add_rPr()
    rPr.set('lang', 'en-US' if spec.get('ltr') else 'he-IL')
    rPr.set('altLang', 'he-IL' if spec.get('ltr') else 'en-US')
    rPr.set('sz', str(max(100, int(round(spec['sz'] * 50 * scale)))))
    rPr.set('b', '1' if spec['b'] else '0')
    rPr.set('dirty', '0')
    if spec.get('sup'):
        rPr.set('baseline', '30000' if spec['sup'] > 0 else '-25000')
    r.font.color.rgb = RGBColor.from_string(spec['c'])
    for tag in ('a:latin', 'a:ea', 'a:cs'):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set('typeface', FONT)


def fill_text(tf, paras, al, rtl, lh, scale=1.0):
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    first = True
    for para in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        pPr = p._p.get_or_add_pPr()
        if rtl:
            pPr.set('rtl', '1')
        p.alignment = ALIGN[al]
        if lh and abs(lh - 1.0) > 0.01:
            p.line_spacing = lh
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        for spec in para:
            style_run(p.add_run(), spec, rtl, scale)


def add_rect(sl, o):
    shp = sl.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if o.get('rad') else MSO_SHAPE.RECTANGLE,
        E(o['x']), E(o['y']), E(o['w']), E(o['h']))
    if o.get('rad'):
        try:
            shp.adjustments[0] = float(o['rad'])
        except Exception:
            pass
    if o.get('fill'):
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string(o['fill'])
    else:
        shp.fill.background()
    if o.get('line'):
        shp.line.color.rgb = RGBColor.from_string(o['line']['c'])
        shp.line.width = Pt(o['line']['w'] * 0.5)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    st = shp._element.find(qn('p:style'))
    if st is not None:
        shp._element.remove(st)
    shp.text_frame.word_wrap = False


def set_cell_borders(cell, colr='FFFFFF', w_pt=1.0):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
        for old in tcPr.findall(qn(tag)):
            tcPr.remove(old)
    # schema order: lnL, lnR, lnT, lnB then fill
    for tag in ('a:lnB', 'a:lnT', 'a:lnR', 'a:lnL'):
        ln = tcPr.makeelement(qn(tag), {'w': str(int(w_pt * 12700)), 'cap': 'flat',
                                        'cmpd': 'sng', 'algn': 'ctr'})
        f = ln.makeelement(qn('a:solidFill'), {})
        c = ln.makeelement(qn('a:srgbClr'), {'val': colr})
        f.append(c)
        ln.append(f)
        tcPr.insert(0, ln)


def add_table(sl, o):
    nr, nc = len(o['rows']), len(o['cols'])
    gf = sl.shapes.add_table(nr, nc, E(o['x']), E(o['y']), E(o['w']), E(o['h']))
    tbl = gf.table
    tblPr = tbl._tbl.find(qn('a:tblPr'))
    tblPr.set('firstRow', '0')
    tblPr.set('bandRow', '0')
    for st in tblPr.findall(qn('a:tableStyleId')):
        tblPr.remove(st)
    sid = tblPr.makeelement(qn('a:tableStyleId'), {})
    sid.text = NOSTYLE
    tblPr.append(sid)
    for i, w in enumerate(o['cols']):
        tbl.columns[i].width = E(w)
    for i, h in enumerate(o['rows']):
        tbl.rows[i].height = E(h)
    for cd in o['cells']:
        r, c = cd['r'], cd['c']
        if r >= nr or c >= nc:
            continue
        cell = tbl.cell(r, c)
        if cd['rs'] > 1 or cd['cs'] > 1:
            try:
                cell.merge(tbl.cell(min(nr - 1, r + cd['rs'] - 1), min(nc - 1, c + cd['cs'] - 1)))
            except Exception:
                pass
    for cd in o['cells']:
        r, c = cd['r'], cd['c']
        if r >= nr or c >= nc:
            continue
        cell = tbl.cell(r, c)
        if cd.get('fill'):
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(cd['fill'])
        else:
            cell.fill.background()
        pt, pr, pb, pl = cd['pad']
        vcap = 0 if cd['va'] == 'ctr' else 4
        cell.margin_top, cell.margin_bottom = E(min(pt, vcap)), E(min(pb, vcap))
        cell.margin_left, cell.margin_right = E(pl), E(pr)
        cell.vertical_anchor = ANCH[cd['va']]
        set_cell_borders(cell)
        fill_text(cell.text_frame, cd['paras'] or [[]], cd['al'], cd['rtl'], cd['lh'],
                  cd.get('sc', 1.0))


IMG = {}
for f in os.listdir(SRC + '/img'):
    IMG[os.path.splitext(f)[0]] = SRC + '/img/' + f

for s in data['slides']:
    sl = prs.slides.add_slide(blank)
    for o in s['ops']:
        if o['t'] == 'rect':
            add_rect(sl, o)
        elif o['t'] == 'pic':
            p = IMG.get(o['img'])
            if p:
                sl.shapes.add_picture(p, E(o['x']), E(o['y']), E(o['w']), E(o['h']))
        elif o['t'] == 'table':
            add_table(sl, o)
        elif o['t'] == 'text':
            fs = max((r['sz'] for para in o['paras'] for r in para), default=16) * float(o.get('sc', 1.0))
            lead = max(0.0, (o['lh'] - 1.0)) * fs / 2.0   # CSS half-leading above line 1
            tb = sl.shapes.add_textbox(E(o['x']), E(o['y'] - lead), E(o['w'] + 2), E(o['h'] + lead + 4))
            tf = tb.text_frame
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            tf.vertical_anchor = MSO_ANCHOR.TOP
            fill_text(tf, o['paras'], o['al'], o['rtl'], o['lh'], float(o.get('sc', 1.0)))

prs.save(DST)
print('saved', DST, len(prs.slides.__iter__.__self__._sldIdLst), 'slides')
